"""
Corrected (R3) versions of the five code-derived "Technical … (Operational)" documents.

Owner instruction 14-Sep-2026: the R2 originals (generated 10-Sep from the code, never product-
authored) contained 4 confirmed-wrong and 46 partly-supported claims. R3 keeps only what the
repository supports, narrows or removes the rest, never describes a backend capability as a UI
step unless the UI element was found, and stays labelled as code-derived. Evidence per claim is
in PROVENANCE.md (written by this script), against repository revision REPO_REV.

  uv run --with python-docx python generated-docs/build_r3.py
    → generated-docs/R3/*.docx (corrected sources), generated-docs/PROVENANCE.md
    R2 originals are copied verbatim into generated-docs/R2/ and are never modified.
"""
from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import docx  # python-docx

HERE = Path(__file__).resolve().parent
R2_SRC = Path("C:/tmp/trag/documents")
REPO_REV = "27a40b2ce"  # feature/chatbot-enterprise, the revision every file:line below was read at
DERIVED = ("Code-derived documentation (revision R3, September 2026): written from the application source code, "
           f"repository revision {REPO_REV}, not by the product team. Where the official manuals cover a task, they take precedence. "
           "Evidence for every statement: central-assistant-py/generated-docs/PROVENANCE.md.")

# (title, audience, intro, [(heading, [paragraphs...])])
DOCS: dict[str, tuple[str, str, str, list[tuple[str, list[str]]]]] = {
    "Technical - Bulk Data Import (Operational) User Manual.docx": (
        "TECHNICAL USER MANUAL — BULK DATA IMPORT (OPERATIONAL)", "Office / Sail Admin",
        "Bulk Data Import loads many records at once from a spreadsheet — Machinery Components, Jobs, Spares, Stores, Locations, "
        "Rotation Items and Work-Order History. The pattern is: download the template → fill it in → upload → the system validates the "
        "file (dry-run) → review the results → import → check the summary.",
        [
            ("1.1.11.2 How To Access Bulk Data Import", [
                "In the PMS module open Admin and select the 'Bulk Data Imp' tab (the page heading reads 'Bulk Data Import').",
                "Select the Vessel from the 'Vessel:' dropdown — every import is associated with the selected vessel.",
                "Choose the template (data type) you want to import from the TEMPLATES list: Machinery Components, Jobs, Spares, Stores, Locations, Rotation Items or History (work-order history).",
            ]),
            ("1.1.11.3 How To Download the Import Template", [
                "For the selected template click 'Download Template'.",
                "Fill in your data keeping the column headings unchanged — rows are read by their exact column names.",
                "Save the file as .csv, .xls or .xlsx (maximum 20 MB).",
                "Vessel Code must match the vessel selected on screen; rows with another vessel code are rejected.",
                "Component Code must already exist for the selected vessel; otherwise the row is rejected with 'Component Code does not exist for the selected vessel'.",
                "Jobs: the Job Code column is optional. A blank Job Code is accepted and a code is generated on import (JOB-XXXXXXX). Duplicate checking applies only to rows that supply a Job Code together with a Component Code for the vessel.",
                "Jobs: 'Assigned To' must be one of the recognised ranks and 'Approver' must be a rank from the Rank Master. The template's Master Data sheet lists the available ranks for reference.",
            ]),
            ("1.1.11.4 How To Upload a File", [
                "Click the upload area ('Click to upload or drag and drop') and select your completed file, or drag it onto the area.",
                "As soon as the file is selected the system validates it (the dry-run). Nothing is saved at this point.",
                "Check that the vessel shown on screen is the one your file was prepared for.",
            ]),
            ("1.1.11.5 Understanding the Dry-Run (Validation) Step", [
                "The dry-run checks the file without saving anything.",
                "Each row is marked ok, warning or error; the summary shows Valid, Warnings, Errors and Total Rows.",
                "Nothing is written to the system until you click the import button.",
            ]),
            ("1.1.11.6 How To Read the Dry-Run Results", [
                "Errors — rows that cannot be imported as they are (for example 'Component Code does not exist for the selected vessel', or a mandatory field is missing). Correct them in the file and upload it again; use 'Re-validate' after changing the file.",
                "Warnings — rows that will import but where a value is adjusted (for example a unit that will be set to Hours).",
                "The counts of created, updated and skipped records are shown after the import runs, not in the dry-run.",
            ]),
            ("1.1.11.7 How To Complete the Import", [
                "When the results are acceptable click 'Import N Records'.",
                "If errors remain you can click 'Skip Errors & Import N Valid Rows' and confirm 'Proceed with Partial Import' — only the valid rows are imported.",
                "The final summary shows how many records were created, updated and skipped; skipped rows carry a reason (for example 'WO already exists').",
            ]),
            ("1.1.11.8 How To Confirm the Import Worked", [
                "Open the relevant screen (for example Components or Spares) for the vessel and confirm the records appear.",
                "Open 'Import History' (button on the Bulk Data Import page, or the Import History tab) to see past imports with their result counts.",
                "To fix rows that were skipped or failed, correct them in your file and upload the file again — there is no selective re-import of individual rows.",
                "An import can be undone as a whole from Import History (Undo), which reverses that import's records.",
            ]),
        ]),
    "Technical - Ship-Side (Vessel Crew) Operational Notes.docx": (
        "TECHNICAL USER MANUAL — SHIP-SIDE (VESSEL CREW) OPERATIONAL NOTES", "Ship (Vessel Crew)",
        "The official user manuals are written 'For Office / Sail Admin', but the same screens and steps apply on the ship. This note covers "
        "only what is different or specific to working on the vessel. For the task steps themselves (components, work orders, running hours, "
        "spares, stores, defects) use the PMS, Defects and Certificates & Surveys manuals — they apply on the ship as well.",
        [
            ("1.1.13.1 How Working on the Ship Differs", [
                "The ship runs its own local copy of the system, so it works without an internet connection.",
                "You do your normal work locally — log running hours, complete work orders, record spares and stores — exactly as described in the Office manuals.",
                "Your changes are exchanged with the office through sync: automatically on the schedule set in the ship's Auto-Sync Settings, or manually with 'Sync Now'.",
                "Jobs and components that the office has deactivated no longer appear to ship users.",
            ]),
            ("1.1.13.2 Completing Work Orders on the Ship — Current Rules", [
                "Running-hours jobs need a reading. When you complete a work order for a job based on running hours you must enter the running-hours reading at completion; the form does not accept the completion without it. Exception: components marked as not running-hours driven do not require a reading. The per-vessel 'RH validation' switch (Sail Admin) can relax the running-hours checks for a vessel.",
                "A reading lower than the component's current hours is refused ('Running hours cannot go backward!'). The one exception is a back-dated entry: if the reading date you enter is earlier than the date of the component's current reading, the completion is saved and the reading is kept on the work order for scheduling, but the component's counter is not changed.",
                "The completion reading cannot be higher than the component's current meter reading (the work cannot have been done at more hours than the meter shows now).",
                "Hours only legitimately go down for a meter replacement or a renewal reset, recorded as such.",
                "Save as Draft: you can save partially filled completion details ('Save Draft') and return later. A draft changes nothing on the live work order and does not change its status until you submit.",
                "Work-order numbers carry the vessel code as a prefix: planned work orders are <VESSELCODE>-<JOBCODE>-<COMPONENTCODE>-<YEAR>-<NNN>, unplanned ones <VESSELCODE>-UWO-<COMPONENTCODE>-<YEAR>-<NNN>. Older work orders keep their original numbers; the running sequence continues across both formats.",
            ]),
            ("1.1.13.3 Approvals as Seen From the Ship", [
                "Work sent for approval (a completed work order awaiting review, a postponement request) is decided on the office side. The ship submits; the office approves or rejects with remarks.",
                "A completion rejected by the office comes back to the vessel with remarks — correct it and resubmit.",
                "A postponed work order stays postponed until someone acts on it — postponements do not revert automatically when their period expires.",
            ]),
            ("1.1.13.4 Day-to-Day Tasks (Use the Existing Manuals)", [
                "The following are identical to the Office manuals — refer to them directly:",
                "Update Running Hours → PMS manual, Running Hours section.",
                "Complete a Work Order → PMS manual, How To Complete The Work Order.",
                "Create an Unplanned Work Order → PMS manual, How To Create An Unplanned Work Order.",
                "Spares / Stores transactions → PMS manual, Spares / Stores sections.",
                "Raise / update a Defect → Defects manual.",
            ]),
            ("1.1.13.5 Sync Tasks on the Ship (See the Sync Manual)", [
                "Change the Auto-Sync schedule — Admin → Sync Dashboard → 'Auto-Sync Settings'.",
                "Run a manual sync — 'Sync Now' on the Sync Dashboard.",
                "Confirm a sync completed — the sync status indicator (connection state, last sync, pending / failed / conflict counts) and the 'Recent Sync History' on the Sync Dashboard.",
            ]),
        ]),
    "Technical - Recent Updates & Changed Behaviours (Operational) Notes.docx": (
        "TECHNICAL USER MANUAL — RECENT UPDATES & CHANGED BEHAVIOURS (OPERATIONAL)", "Office · Sail Admin · Ship (Vessel Crew)",
        "The PMS, Defects and Certificates & Surveys user manuals were last revised in June 2026. The system has changed since. This note lists "
        "the behaviours that are new or different from what those manuals describe, in plain user terms.",
        [
            ("1.1.14.1 Work-Order Numbers Now Include the Vessel Code", [
                "New work orders are numbered with the vessel's code as a prefix: planned work orders <VESSELCODE>-<JOBCODE>-<COMPONENTCODE>-<YEAR>-<NNN>, unplanned work orders <VESSELCODE>-UWO-<COMPONENTCODE>-<YEAR>-<NNN>.",
                "Existing work orders keep their original numbers; the running sequence continues across old and new formats, so numbering does not restart.",
                "If a vessel's code is not present on an instance, its new work orders use the older, non-prefixed format. The vessel code reaches a ship only through provisioning — the vessel master is not part of the routine sync — so a ship still producing old-format numbers needs its vessel code applied by re-provisioning or by support; the office 'Sync All' master refresh does not do this.",
            ]),
            ("1.1.14.2 Completing a Running-Hours Work Order Requires a Reading", [
                "When completing a work order for a job based on running hours, the running-hours reading at completion is mandatory — the completion is refused without it. Components marked as not running-hours driven are exempt.",
                "The reading is recorded on the work order together with the outcome of applying it (whether it updated the component's counter, or was kept on record without changing the counter — see the next item).",
            ]),
            ("1.1.14.3 Running Hours Never Move Backwards", [
                "A component's running-hours counter is never reduced by a late, low or out-of-order reading from any source (ship entry, office entry, work-order completion or sync). A lower reading entered directly is refused; a lower reading that is back-dated is kept in the audit history and on the work order without changing the counter.",
                "The only legitimate decreases are a meter replacement or a renewal reset, recorded as such when entered.",
                "Across sync, readings are reconciled by reading date — the latest date wins; a same-date tie goes to the ship. See the Sync (Operational) manual, section on running hours.",
            ]),
            ("1.1.14.4 Approving Work Orders Requires an Explicit Decision", [
                "A work order in Pending Approval moves on only when the approver explicitly chooses Approve or Reject. Its status no longer changes out of Pending Approval as a side effect of editing the record.",
            ]),
            ("1.1.14.5 Save as Draft on Work-Order Completion", [
                "Completion details can be saved as a draft ('Save Draft') and finished later. A draft changes nothing on the live work order until it is submitted, and the work order's status stays as it was.",
            ]),
            ("1.1.14.6 Office Generation of Work Orders Is a Per-Vessel Switch", [
                "Only a Sail Admin can generate work orders directly from the office, and only for a vessel whose 'office work-order generation' switch is enabled (off by default). The switch is on the 'Lead Time & Grace Period Settings' screen.",
                "With the switch off, the Sail Admin is told that office work-order generation is not enabled for the vessel and where to enable it. Other office roles are refused with a role message.",
                "Ships always generate their own work orders on their daily scan, regardless of this switch.",
                "The office 'Generate Now' button (Work Orders screen, with one vessel selected) runs an immediate, vessel-scoped generation instead of waiting for the daily sweep, subject to the same switch and role rule.",
            ]),
            ("1.1.14.7 Office Entry of Running Hours Is a Per-Vessel Switch", [
                "Office users can enter running hours for a vessel only if that vessel's 'office running-hours entry' switch is enabled (off by default). An office entry must carry an explicit reading date. Ship entry is unaffected.",
            ]),
            ("1.1.14.8 Superintendent Capabilities", [
                "Office roles (Office, PMS Admin, Sail Admin) can reopen a completed work order with remarks and can reject a completion back to the vessel for correction. A rejected work order returns to the vessel's normal flow and appears in the Due or Overdue tab according to its due date.",
                "The Superintendent page groups activity for the currently selected vessel into pending, acknowledged and information tabs.",
                "A per-vessel superintendent lock switch (Sail Admin, 'Lead Time & Grace Period Settings' screen) restricts certain vessel actions to the superintendent when enabled.",
            ]),
            ("1.1.14.9 Postponements Do Not Auto-Revert", [
                "A postponed work order stays postponed when its postponement period expires — it returns to the normal flow only when a user acts on it. This is intentional; do not wait for an automatic status change.",
            ]),
            ("1.1.14.10 Rotational (Swappable) Parts Keep Their Own Identity", [
                "Rotational items (parts swapped between positions, for example exchanged pumps or motors) are tracked by a stamp identity that follows the physical part. When parts are swapped, running hours follow the part, not the position, and the rotation history is kept.",
            ]),
            ("1.1.14.11 Certificates: Next Annual / Next Interim Dates", [
                "Vessel certificate records carry Next Annual and Next Interim date fields, entered per vessel on the certificate record. Existing records show a dash for these until a user fills them in.",
            ]),
            ("1.1.14.12 Change-Request and Postponement Approval Levels", [
                "Change requests and work-order postponement requests go through up to two approval levels (Level 1, Level 2), each enabled per function under Admin → Approval Workflow. The approvers for each level are maintained separately from that screen. Each decision (approve / reject, with comments) is recorded against the request.",
            ]),
        ]),
    "Technical - Roles & Permissions (Operational) User Manual.docx": (
        "TECHNICAL USER MANUAL — ROLES & PERMISSIONS (OPERATIONAL)", "Office / Sail Admin",
        "Different users see and do different things depending on their role and their data scope (which vessels they can see). This note "
        "explains what each role can do and how scope works.",
        [
            ("1.1.12.1 Overview — Roles at a Glance", [
                "Sail Admin — the highest-level administrator; full access across the module, including Admin functions (bulk import, configuration, approval workflow setup).",
                "Superintendent — office-side user managing maintenance across assigned vessels; can review, approve and export. (Users with the 'Admin' profile role are shown as Superintendent.)",
                "Client Admin — the default role for office users; admin-level scope (sees all vessels).",
                "Head of Dept — department head; can use the 'Me' / 'My Team' toggle on the PMS Dashboard. Users whose profile role is 'Vessel Admin' are treated as Head of Dept in this module — this is the 'Vessel Admin' the PMS manual refers to when it says only Vessel Admin users can use the 'My Team' toggle.",
                "Vessel — ship crew; works on the vessel's own data day-to-day (running hours, completing work orders, spares/stores).",
                "External — external user role with restricted access.",
                "Level 2 Reviewer (PMS manual §1.1.3.5 'How to Approve Work Orders (Level 2 Reviewer Role)') — assigned per job by rank: the job form's 'Level 2 Reviewer (Rank)' field decides which rank reviews that job's work orders at the 'Pending Office Review' stage. It is not a role a Sail Admin assigns to a user; the reviewer actions themselves are available to the Office, PMS Admin and Sail Admin roles.",
            ]),
            ("1.1.12.2 Understanding Data Scope (My Vessel / My Team / Me)", [
                "Scope controls which vessels' data a user sees. When an office user selects 'My Vessel' as the scope, only the vessels assigned to that user are shown.",
                "The 'Me' / 'My Team' toggle exists on the PMS Dashboard (Operation tab, in the header) only — see the PMS manual §1.1.3.4. It is shown to users who are not in admin scope (i.e. not Sail Admin, Client Admin or Superintendent) — in practice Head of Dept users, including those with the 'Vessel Admin' profile role — it filters the work-order figures only (spares, anomalies and PMS requests stay vessel-wide), and it is disabled until the user's team hierarchy is mapped.",
                "Ship users see only active jobs and components: jobs and components deactivated by the office no longer appear on the vessel.",
                "If you cannot see a vessel, a record, or a colleague's work, it is usually because of your scope or role, not a fault.",
            ]),
            ("1.1.12.3 How Permissions Are Enforced", [
                "Certain actions are restricted by role on the server, for example maintaining fleet master list types (Sail Admin only), reviewer actions on work orders (Office, PMS Admin, Sail Admin), and office work-order generation (Sail Admin). A refused action returns the message 'Forbidden - Insufficient permissions'.",
                "Menu access per role is configured by the Sail Admin and controls what is shown in the navigation.",
                "Seeing a 'Forbidden' message means your role does not allow the action — speak to your Sail Admin; it is not a system fault.",
            ]),
            ("1.1.12.4 What a Sail Admin Can Do", [
                "Full access across PMS, Defects, Certificates & Surveys.",
                "Access the Admin area, including Bulk Data Import, role/menu configuration, Approval Workflow setup and the per-vessel settings ('Lead Time & Grace Period Settings').",
                "See all vessels.",
            ]),
            ("1.1.12.5 What a Superintendent (Office) User Can Do", [
                "View and manage maintenance for assigned vessels (use the 'My Vessel' scope to focus on them).",
                "Review dashboards, work orders, running hours, spares, stores, defects, certificates and surveys.",
                "Review postponement requests and approve / reject work as permitted by the approval configuration.",
                "Reopen a completed work order or reject a completion back to the vessel with remarks (available to the Office, PMS Admin and Sail Admin roles), and follow the selected vessel's activity on the Superintendent page (pending / acknowledged / information tabs).",
                "Export records from the screens that offer an Export button (dashboard, work orders, components, running hours, anomalies, planner and others).",
            ]),
            ("1.1.12.6 What a Head of Dept Can Do", [
                "Use the office screens for the assigned vessels, and switch between 'Me' (own work) and 'My Team' (team's work) on the PMS Dashboard.",
                "Head of Dept does not approve postponement requests when no approval steps are configured — that is reserved for office roles.",
            ]),
            ("1.1.12.7 What a Ship / Vessel User (Crew) Can Do", [
                "Work on the vessel's own local system day-to-day: log running hours, complete work orders, record spares/stores consumption.",
                "Trigger a manual 'Sync Now' and manage the Auto-Sync Settings on the vessel's Sync Dashboard (see the Sync (Operational) manual).",
            ]),
            ("1.1.12.8 Approval Workflows (Admin)", [
                "Change-request and postponement approvals use up to two levels (Level 1, Level 2), enabled per function under Admin → Approval Workflow. The approvers for each level are maintained separately from that screen. Each decision (approve / reject, with comments) is recorded against the request.",
            ]),
            ("1.1.12.9 Per-Vessel Administrative Switches (Sail Admin)", [
                "Some behaviours are switched per vessel by a Sail Admin on the 'Lead Time & Grace Period Settings' screen:",
                "Office work-order generation — whether a Sail Admin may generate this vessel's planned work orders from the office (off by default; ships always generate their own regardless).",
                "Office running-hours entry — whether office users may enter running hours for this vessel (off by default; an explicit reading date is required).",
                "RH validation — whether the running-hours checks are applied for this vessel (on by default).",
                "Superintendent lock — restricts certain vessel actions to the superintendent when enabled.",
                "If an action is unavailable for one vessel but works for another, check these per-vessel switches before suspecting a permissions problem.",
            ]),
        ]),
    "Technical - Sync (Operational) User Manual.docx": (
        "TECHNICAL USER MANUAL — SYNC (OPERATIONAL)", "Office · Sail Admin · Ship (Vessel Crew)",
        "Ships at sea run their own local copy of the system and exchange data with the office when a connection is available. This note "
        "explains how to see that a vessel is syncing, what the status means, how to sync manually, and what happens when the ship and the "
        "office change the same record.",
        [
            ("1.2.1 How To View Vessel Sync Activity (Office)", [
                "Audience: Office / Sail Admin (shore only).",
                "Go to Admin → Fleet Overview (the 'Fleet Sync Overview' screen).",
                "For each vessel it shows the last sync time (for example '5h ago', '2d ago', or 'Never') and a colour-coded status; the list refreshes automatically every 60 seconds.",
                "The 'Recent Fleet Activity' table lists recent syncs with records sent and received, conflicts and duration.",
                "The eye icon on a vessel row opens the Sync Dashboard; select the vessel there (it is pre-selected only when a single vessel exists) to see its history — date/time, status, pushed, pulled, conflicts, files and duration per sync.",
            ]),
            ("1.2.2 Understanding Sync Status Colours", [
                "Audience: Office / Sail Admin.",
                "Green (Synced) — the vessel synced within the last 24 hours and has no unresolved conflicts.",
                "Amber (Stale) — no sync for more than 24 hours. This is time-based only; changes waiting to go out are shown as a separate 'pending changes' count, not as a colour.",
                "Red (Overdue / Conflicts) — unresolved data conflicts, or no sync for more than 48 hours.",
                "Grey (Never Synced / Not Provisioned) — the vessel is newly set up or has not completed its first sync.",
                "Note: the 'Stale' summary card at the top counts vessels with no sync for more than 48 hours, a different threshold from the amber badge (24 hours).",
                "The list is ordered by oldest last-sync first (never-synced vessels first), so long-quiet vessels are at the top. A vessel that synced recently but has conflicts is not moved up — look for the red badge.",
            ]),
            ("1.2.3 How To Confirm a Vessel Is Syncing Daily", [
                "Open the Fleet Sync Overview (Admin → Fleet Overview).",
                "Confirm the vessel shows green with a recent last-sync time.",
                "Vessels with no recent sync are at the top of the list in amber or red — those are the ones to follow up.",
            ]),
            ("1.2.4 Auto-Sync vs Manual Sync (What's Normal)", [
                "Audience: Office / Sail Admin + Ship.",
                "Sync is always started by the ship; the office has no sync trigger. Auto-sync runs on the ship when enabled in the ship's Auto-Sync Settings (enabled by default, every 6 hours, first run about 3 minutes after the ship system starts).",
                "A manual 'Sync Now' is available on the ship for an immediate sync — for example after a large data entry, or to confirm connectivity once a vessel is back online.",
                "'Sync Now' repeats push/pull cycles until the backlog is clear, or until it reaches 20 cycles, about 60 seconds, or makes no further progress; the remainder completes on the next 'Sync Now' or the next automatic run. After a long outage more than one press may be needed.",
                "Recommended practice: leave auto-sync on for normal daily operation and use 'Sync Now' only when an immediate update is required.",
            ]),
            ("1.2.5 How To Change the Auto-Sync Schedule (Ship)", [
                "Audience: Ship (Vessel Crew) — set on the ship, per ship instance (one vessel per instance).",
                "On the vessel go to Admin → Sync Dashboard and open the 'Auto-Sync Settings' card.",
                "Switch auto-sync on or off and set the 'Sync Interval (minutes)'. A changed interval takes effect immediately without a restart; the next run happens one interval after the change.",
            ]),
            ("1.2.6 How To Run a Manual Sync ('Sync Now') (Ship)", [
                "Audience: Ship (Vessel Crew).",
                "On the vessel's Sync Dashboard click 'Sync Now'.",
                "The ship sends its changes to the office and receives the office's changes; the result (records pushed, records pulled, conflicts, files queued, duration) is shown and a history entry is recorded.",
            ]),
            ("1.2.7 How To Confirm a Sync Completed (Ship & Office)", [
                "On the ship: the sync status indicator shows the connection state (Online / Offline), the last sync time and the pending / failed / conflict counts, with a link to the Sync Dashboard; the dashboard's 'Recent Sync History' lists each sync.",
                "On the office side: the Fleet Sync Overview shows last sync time and status per vessel, with the Sync Dashboard for a vessel's full history.",
            ]),
            ("1.2.8 What Happens When Ship and Office Edit the Same Record", [
                "Audience: Office / Sail Admin.",
                "Ship and office changes to different fields of the same record both survive — the system merges them field by field.",
                "If both sides change the same field of the same record, the more recent change wins and the losing change is recorded as a conflict, not lost. When an incoming change is rejected on the office side, an in-app notification is recorded for the user who made the rejected change (normally the ship user); no notification is raised on the ship side.",
                "Special case — both sides complete the same work order while the vessel was disconnected: this is detected as a conflict and shown on the 'Sync Conflict Review' screen (menu: Conflict Review) with both completions side by side. An office user resolves it by choosing which completion stands; the record is resolved as one unit, never mixed field by field. Until it is resolved the system shows an interim result: the later completion date, or the ship's completion when the dates are the same.",
                "Vessels with unresolved conflicts show red on the Fleet Sync Overview.",
            ]),
            ("1.2.9 How Running Hours Behave Across Sync", [
                "Audience: Office / Sail Admin + Ship.",
                "Running-hours readings from ship and office are reconciled by reading date — the reading with the latest date wins, no matter which side entered it or which arrived first. If two readings carry the same date, the ship's reading wins.",
                "A component's hours never move backwards because of sync. If an older or lower reading arrives late, it is kept in the audit history but does not reduce the counter on screen.",
                "The only cases where hours legitimately go down are a meter replacement or a renewal reset, recorded as such when the reading is entered.",
                "For support: if a reading looks wrong after a sync, check the reading dates in the audit history first — the latest date wins, not the latest arrival.",
            ]),
            ("1.2.10 Vessel Master Data (Office)", [
                "Audience: Office / Sail Admin.",
                "Vessel master records (including the vessel code used as the prefix of work-order numbers) are maintained on the office side and refreshed with the 'Sync All' button on Admin → Data Masters (shore only), which pulls from the company's master system.",
                "Vessel master data reaches a ship when the ship is provisioned; it is not part of the routine data sync. If a vessel's code or master details change, a Sail Admin re-provisions the ship from Admin → Ship Provisioning, or contacts support.",
            ]),
        ]),
}

# claim-level provenance: (document, claim in R3 or removed R2 claim, disposition, evidence)
PROVENANCE: list[tuple[str, str, str, str]] = [
    ("Bulk Data Import", "Importable types = Machinery Components / Jobs / Spares / Stores / Locations / Rotation Items / History (no 'Work Orders')", "CORRECTED", "client/src/pages/admin/BulkDataImport.tsx:266-274"),
    ("Bulk Data Import", "Path = PMS → Admin → 'Bulk Data Imp' tab, heading 'Bulk Data Import'", "CORRECTED", "client/src/components/SideMenuBar.tsx:69; client/src/pages/pms/PMSAdmin.tsx:41,75; client/src/pages/TechnicalModule.tsx:228-229"),
    ("Bulk Data Import", "Vessel dropdown; imports associated with the selected vessel", "KEPT", "client/src/pages/admin/BulkDataImport.tsx:297-309,393"),
    ("Bulk Data Import", "'Download Template' button; headings must stay unchanged; .csv/.xls/.xlsx ≤ 20 MB", "KEPT", "UniformBulkUpload.tsx:758-761,330-338,905-907; server/modules/bulk-upload/routes.ts:14-17,25; validationService.ts:671-680"),
    ("Bulk Data Import", "Vessel Code must match; Component Code must exist", "KEPT", "server/modules/bulk-upload/services/validationService.ts:658-666,995,1625"),
    ("Bulk Data Import", "R2: 'Job codes must be filled in and unique — blank or duplicate rejected'", "REMOVED (wrong) → Job Code optional, auto-generated; duplicates on Job Code + Component Code + Vessel Code", "validationService.ts:217,441-443; client/src/pages/admin/JobUpload.tsx:50"),
    ("Bulk Data Import", "Jobs: 'Assigned To' recognised ranks; 'Approver' from Rank Master; Master Data sheet for reference", "NARROWED", "validationService.ts:1762-1771,448-469,1775-1792; templateService.ts:289-316"),
    ("Bulk Data Import", "Upload area 'Click to upload or drag and drop'; dry-run runs automatically on selection", "CORRECTED", "UniformBulkUpload.tsx:902-907,342-349,357,382"),
    ("Bulk Data Import", "Dry-run rows ok/warning/error; summary Valid/Warnings/Errors/Total; created/updated/skipped only after import", "CORRECTED", "UniformBulkUpload.tsx:79-84,1057-1059,978-1007,475; ImportProgressOverlay.tsx:158-161"),
    ("Bulk Data Import", "'Re-validate' after correcting; warning example (unit set to Hours)", "KEPT", "UniformBulkUpload.tsx:930-936; validationService.ts:1717-1718"),
    ("Bulk Data Import", "Buttons 'Import N Records' / 'Skip Errors & Import N Valid Rows' / 'Proceed with Partial Import'", "CORRECTED", "UniformBulkUpload.tsx:1176,1187,1394"),
    ("Bulk Data Import", "Skipped reasons e.g. 'WO already exists'", "KEPT", "server/modules/bulk-upload/services/importService.ts:1596,1614"),
    ("Bulk Data Import", "'Import History' (not 'Bulk Import History')", "CORRECTED", "BulkDataImport.tsx:402; UniformBulkUpload.tsx:774-776,1241; routes.ts:35"),
    ("Bulk Data Import", "No selective re-import; whole-import Undo from Import History", "CORRECTED (Undo added — UI dialog exists)", "UniformBulkUpload.tsx:1385,200-202,640-642; routes.ts:41"),
    ("Ship-Side", "Only jobs and components deactivated by the office disappear for ship users (spares claim removed — no ship-viewer filter found)", "NARROWED", "server/modules/jobs/controllers/jobController.ts:9-36; server/modules/components/controllers/componentController.ts:13-20; no match in server/modules/spares"),
    ("Ship-Side", "RH reading mandatory; NOT_RH_DRIVEN exempt; per-vessel RH validation switch", "NARROWED", "workOrderCompletionService.ts:131-141; shared/schema.ts:2379-2381; client/src/pages/pms/rhValidationPolicy.ts:11-18"),
    ("Ship-Side", "R2: 'a lower reading is still saved, counter not reduced'", "CORRECTED → lower reading refused; back-dated exception; completion RH ≤ current reading", "server/modules/running-hours/services/rhTimelineValidationService.ts:310-311; workOrderCompletionService.ts:146-156,257-265,358-366"),
    ("Ship-Side", "WO number format with vessel code prefix (JOBCODE segment, UWO for unplanned)", "CORRECTED", "server/utils/workOrderNumbering.ts:7-11,60-86,119-131"),
    ("Ship-Side", "Save Draft; approvals from office; rejected completion returns with remarks; no postponement auto-revert; sync tasks", "KEPT", "WorkOrderFormPage.tsx:2841-2850,7453; workOrderService.ts:2860-2879; server/routes.ts:687-696; SyncDashboard.tsx:525-551,641-647"),
    ("Ship-Side", "Indicator shows connection, last sync, pending/failed/conflict counts; history on the dashboard", "NARROWED", "client/src/components/SyncStatusIndicator.tsx:63-139; SyncDashboard.tsx:659-714"),
    ("Recent Updates", "R2: 'Admin: run Sync Masters to populate vessel codes'", "REMOVED (wrong) → v_code reaches ships via provisioning only; 'Sync All' does not", "docs/SYNC-ARCHITECTURE.md:24-27; shared/syncConfig.ts:1191-1202; client/src/pages/admin/DataMasters.tsx:391,558-575"),
    ("Recent Updates", "WO numbering, legacy fallback, sequence continuity", "KEPT", "server/utils/workOrderNumbering.ts:7-11,66-81,112-126,156-164"),
    ("Recent Updates", "RH mandatory with NOT_RH_DRIVEN exemption; outcome recorded", "NARROWED", "workOrderCompletionService.ts:131-142; shared/schema.ts:1399-1403"),
    ("Recent Updates", "RH never decreases; lower direct entry refused; back-dated kept; reset exceptions; date-wins across sync", "NARROWED", "rhValidation.ts:44-82; rhTimelineValidationService.ts:310-311; docs/SYNC-ARCHITECTURE.md:17-33"),
    ("Recent Updates", "Explicit approve/reject; Save Draft; per-vessel office generation switch; ships ungated; Generate Now", "KEPT", "workOrderService.ts:1187-1191; shared/schema.ts:1405-1409,2371-2373; workOrderGenerationGate.ts:64-87,119-120,150-162; WorkOrders.tsx:1492-1498"),
    ("Recent Updates", "Only a Sail Admin may run office generation; message names 'Lead Time & Grace Period Settings'; other roles refused", "CORRECTED", "workOrderGenerationGate.ts:44,141-148,153-161; PmsVesselSettingsManagement.tsx:402,505-511"),
    ("Recent Updates", "Office RH entry switch + explicit reading date required", "NARROWED", "shared/schema.ts:2375-2377; workOrderCompletionService.ts:192-201,211-218"),
    ("Recent Updates", "Reopen/reject available to Office, PMS Admin, Sail Admin; rejected WO re-computes to Due or Overdue", "CORRECTED", "work-orders/routes.ts:125-135; workOrderController.ts:636; shared/workOrders/status.ts:253; workOrderBulkService.ts:369,381-382"),
    ("Recent Updates", "Superintendent page scoped to the selected vessel", "CORRECTED", "client/src/pages/pms/SuperintendentPage.tsx:53-79,367-379"),
    ("Recent Updates", "Superintendent lock switch; postponements no auto-revert; rotational stamp identity", "KEPT", "shared/schema.ts:2383-2385,375-380,408-422; server/routes.ts:687-694"),
    ("Recent Updates", "Certificates Next Annual / Next Interim; empty shows a dash", "NARROWED", "shared/schema.ts:3506,3508; client/src/pages/cert-surveys/CertificatesPage.tsx:775,812-813,1001-1003"),
    ("Recent Updates", "Two approval levels enabled per function; approvers maintained separately", "CORRECTED", "client/src/pages/admin/ApprovalWorkflow.tsx:148,155-156,197-198,626-629; server/routes.ts:361-369; server/postgresStorage.ts:5785,5946-5950"),
    ("Roles", "Role list = Sail Admin, Superintendent, Client Admin, Head of Dept, Vessel, External; 'Admin' → Superintendent; 'Vessel Admin' → Head of Dept", "CORRECTED", "shared/uiRoles.ts:1-33"),
    ("Roles", "R2: 'Level 2 Reviewer' and 'Vessel Admin' as assignable roles", "CORRECTED (R3.1) → Level 2 Reviewer = per-job rank field (PMS manual §1.1.3.5 uses the term); 'Vessel Admin' = profile role treated as Head of Dept (PMS manual §1.1.3.4 uses the term)", "shared/uiRoles.ts:32; shared/schema.ts:1204,2578; client/src/pages/pms/JobsFormPage.tsx:1346; workOrderService.ts:1720-1724; work-orders/routes.ts:113-123; workOrderBulkService.ts:143-156,227; official PMS Office manual p11-12"),
    ("Roles", "My Vessel scope", "KEPT", "client/src/contexts/VesselContext.tsx:13-60; AuthContext.tsx:447"),
    ("Roles", "Me/My Team: PMS Dashboard Operation tab only; shown when not admin scope; WO KPIs only; disabled until hierarchy mapped", "CORRECTED", "client/src/pages/pms/Dashboard.tsx:477,529-534,636,2836-2876"),
    ("Roles", "R2: 'only HOD' / 'only Vessel Admin' can use the toggle", "REMOVED (contradictory)", "Dashboard.tsx:2836 (`!isAdminScope`)"),
    ("Roles", "Ship users see only active jobs and components", "NARROWED", "jobController.ts:9-36; componentController.ts:13-20"),
    ("Roles", "R2: 'write actions are permission-checked on the server across the module'", "CORRECTED → role checks on specific routes; menu access configuration controls navigation", "server/middleware/permissions.ts:16-27,93-127; server/middleware/auth.ts:37,79; fleet/routes.ts:8-18; work-orders/routes.ts:113-135"),
    ("Roles", "Sail Admin full access; admin area incl. per-vessel settings screen name", "KEPT/NARROWED", "auth.ts:37; permissions.ts:108; PMSAdmin.tsx:41; SideMenuBar.tsx:147; PmsVesselSettingsManagement.tsx:402"),
    ("Roles", "Export on screens that offer it (not 'each screen')", "NARROWED", "Export buttons in 10 PMS pages (Dashboard, WorkOrders, Components, RunningHours, AnomaliesPage, WorkOrderPlanner, …)"),
    ("Roles", "R2: 'HOD can do everything an office user can for their department'", "REMOVED (wrong) → Head of Dept: Me/My Team; excluded from postponement approval when no steps configured", "client/src/components/PostponeApprovalDialog.tsx:109-113; UIRoleContext.tsx:85; ModifyPMS.tsx:178"),
    ("Roles", "Crew: RH, WOs, spares/stores; Sync Now; Auto-Sync Settings", "KEPT", "workOrderCompletionService.ts:185-201; SyncDashboard.tsx:525-531,641-647; AutoSyncSettingsCard.tsx:98"),
    ("Roles", "Per-vessel switches on 'Lead Time & Grace Period Settings' incl. RH validation (default on)", "CORRECTED", "PmsVesselSettingsManagement.tsx:402,617; shared/schema.ts:2371-2385"),
    ("Sync", "Fleet Overview path, screen name, relative times, colours, 60 s auto-refresh", "KEPT", "SideMenuBar.tsx:142-143; SyncFleetOverview.tsx:7,87-148,199,221,313,407,441-448"),
    ("Sync", "Eye icon opens the Sync Dashboard without passing the vessel; vessel pre-selected only when one exists; activity table has no Files column", "CORRECTED", "SyncFleetOverview.tsx:477-488,707-714; SyncDashboard.tsx:211-215,439-452,684-690"),
    ("Sync", "R2: 'amber also means changes waiting to go out'", "REMOVED (wrong) → stale is time-based only; pending changes separate count", "SyncFleetOverview.tsx:134-142,450-457"),
    ("Sync", "Stale card counts >48 h vs badge >24 h (in-code inconsistency, stated)", "ADDED", "SyncFleetOverview.tsx:136-142,257-262,362"),
    ("Sync", "Sorting = oldest last-sync first, never-synced first; conflicts do not move up", "CORRECTED", "server/modules/sync/service.ts:1640"),
    ("Sync", "Sync is ship-initiated; auto-sync ship-only when enabled (default true since migration 119), 6 h, 3-min boot delay", "CORRECTED", "autoSyncScheduler.ts:26-27,263-279; migrations/103_sync_settings_seed.sql:10; migrations/119_sync_settings_defaults_update.sql:8-24; SyncFleetOverview.tsx:242-243"),
    ("Sync", "Sync Now stops at 20 cycles / ~60 s / no progress", "CORRECTED", "server/modules/sync/syncEngine.ts:62-63,568,595-611; controller.ts:302-305"),
    ("Sync", "Auto-Sync Settings per ship instance; interval effective immediately, next run after one interval", "NARROWED", "AutoSyncSettingsCard.tsx:20,98,112,127-128; autoSyncScheduler.ts:160-193,285-288; sync/controller.ts:686-690"),
    ("Sync", "Indicator content (Online/Offline, last sync, counts, link); history on dashboard", "NARROWED", "SyncStatusIndicator.tsx:63-139; SyncDashboard.tsx:659-714"),
    ("Sync", "Field-level merge; same-field: recent wins, conflict recorded; notification to the rejected sender on the office side; none on ship", "CORRECTED", "shared/syncConfig.ts:6; oneWayApplier.ts:1622-1635; sync/service.ts:411-447"),
    ("Sync", "Dual completion → Sync Conflict Review (menu 'Conflict Review'); one unit; interim later date, tie → ship; red on overview", "KEPT", "dualCompletionResolver.ts:10-15,89-103; SyncConflictReview.tsx:249,347; SideMenuBar.tsx:97; conflictReviewRepository.ts:742-745,896-899; SyncFleetOverview.tsx:120-125"),
    ("Sync", "RH across sync: latest date wins, tie → ship, never backwards, reset exceptions", "KEPT", "rhEventComparator.ts:4-17,77-121,200-210,258; oneWayApplier.ts:1077-1086"),
    ("Sync", "'Sync All' on Admin → Data Masters (shore only); vessel code is the WO prefix", "CORRECTED (label) / KEPT (WO prefix, hand-verified)", "DataMasters.tsx:558-580; server/modules/misc/routes.ts:32; workOrderNumbering.ts:7-11,60-86"),
    ("Sync", "Vessel master at provisioning only; re-provision from Admin → Ship Provisioning", "CORRECTED", "shared/syncConfig.ts:1191-1202; SideMenuBar.tsx:98; provisioningService.ts:55,373,851"),
]


def build(name: str, spec: tuple[str, str, str, list[tuple[str, list[str]]]], out: Path) -> None:
    title, audience, intro, sections = spec
    d = docx.Document()
    d.add_heading(title, level=1)
    d.add_paragraph(f"Revision: R3 (September 2026) / Audience: {audience}")
    d.add_paragraph(DERIVED)
    d.add_paragraph(intro)
    for heading, paras in sections:
        d.add_heading(heading, level=2)
        for p in paras:
            d.add_paragraph(p, style="List Bullet")
    d.save(out)


def main() -> None:
    import sys
    only = set(sys.argv[1:])  # optional: rebuild only the named documents (others keep their verified bytes)
    (HERE / "R2").mkdir(exist_ok=True)
    (HERE / "R3").mkdir(exist_ok=True)
    rows = ["# Provenance of the code-derived Technical documents (R3, 14-Sep-2026)", "",
            f"Repository revision read: `{REPO_REV}` (branch feature/chatbot-enterprise). R2 = the 10-Sep originals, preserved unchanged in `R2/`; "
            "R3 = corrected sources in `R3/` built by `build_r3.py`. Dispositions: KEPT (evidence supports the R2 wording), NARROWED (wording reduced to what "
            "the evidence supports), CORRECTED (wording replaced), REMOVED (claim dropped), ADDED (fact the code establishes that R2 lacked).", "",
            "| document | sha256 R2 | sha256 R3 |", "|---|---|---|"]
    for name, spec in DOCS.items():
        src = R2_SRC / name
        r2 = HERE / "R2" / name
        if not r2.exists():
            shutil.copy2(src, r2)
        r3 = HERE / "R3" / name
        if not only or name in only or not r3.exists():
            build(name, spec, r3)
        rows.append(f"| {name} | {hashlib.sha256(r2.read_bytes()).hexdigest()[:16]} | {hashlib.sha256(r3.read_bytes()).hexdigest()[:16]} |")
        print(f"built {r3.name}: {r3.stat().st_size} bytes")
    rows += ["", "| document | claim (R3 wording, or the R2 claim removed) | disposition | evidence (file:line at the revision above) |", "|---|---|---|---|"]
    rows += [f"| {d} | {c} | {s} | {e} |" for d, c, s, e in PROVENANCE]
    (HERE / "PROVENANCE.md").write_text("\n".join(rows) + "\n", encoding="utf-8")
    print("PROVENANCE.md written,", len(PROVENANCE), "claims")


if __name__ == "__main__":
    main()
