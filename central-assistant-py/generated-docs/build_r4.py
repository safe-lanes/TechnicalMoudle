# Build R4 of the Recent Updates note from the revision that is ACTUALLY INDEXED in the candidate set kb-pilot-c
# (sha256 d10f3027…, recovered from git blob 619dbb92 at commit 972f424d7 and kept beside this script as
# "...AS-INDEXED-kb-pilot-c.docx"). Only §1.1.14.6 is touched: its heading is retitled and its four bullets are
# replaced by six, one per action, each carrying its code reference. Every other paragraph is left byte-identical.
#   python generated-docs/build_r4.py
import copy
import os

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "R3", "Technical - Recent Updates & Changed Behaviours (Operational) Notes.AS-INDEXED-kb-pilot-c.docx")
OUT = os.path.join(HERE, "R4", "Technical - Recent Updates & Changed Behaviours (Operational) Notes.docx")

HEADING = "1.1.14.6 Office Generation of Work Orders — Conditions by Action"
BULLETS = [
    "Office 'Generate Now' (Work Orders screen, one vessel selected) — runs an immediate, vessel-scoped generation "
    "instead of waiting for the daily sweep. Two conditions apply, both on the server: the caller's role must be Sail "
    "Admin, and that vessel's 'office work-order generation' switch must be enabled. The switch is off by default and "
    "is set on the 'Lead Time & Grace Period Settings' screen. [code: workOrderController.generateNow → "
    "workOrderGenerationGate.evaluateDirectGeneration; PROVISIONING_ROLE = 'Sail Admin'; isOfficeWoGenerationEnabled]",

    "Refusals for 'Generate Now' are distinct. A caller who is not a Sail Admin is refused with: \"Only a Sail Admin "
    "may generate work orders directly from the office.\" A Sail Admin acting on a vessel whose switch is off is "
    "refused with: \"Office work-order generation is not enabled for this vessel. A Sail Admin can enable it per "
    "vessel on the Lead Time & Grace Period Settings screen.\" A vessel whose provisioning state cannot be read is "
    "refused as well — the check fails closed. [code: evaluateDirectGeneration, codes ROLE_NOT_PERMITTED, "
    "OFFICE_GENERATION_DISABLED, VESSEL_STATE_UNKNOWN]",

    "Office per-job 'Generate WO' (Components → the component's job, choose a reason) — the vessel's 'office "
    "work-order generation' switch must be enabled, and the job must exist, be active and not already have an active "
    "work order; the reason must be Planning, Breakdown or Other. There is NO action-specific role check on this "
    "path: it does not go through the role gate. Any signed-in user with access to the vessel may use it while the "
    "switch is on. [code: jobService.generateWorkOrder → isOfficeWoGenerationEnabled only; "
    "jobDueScanner.generateWorkOrderForJob for the active-work-order check]",

    "Unplanned work order ('+ Unplanned W.O' on the Work Orders screen) — neither the Sail Admin rule nor the office "
    "generation switch applies to it. Normal sign-in and vessel access still apply. [code: the create path does not "
    "call the generation gate]",

    "Ships always generate their own work orders on their daily scan, regardless of this switch; a ship instance is "
    "allowed through the gate without the role or switch test. [code: evaluateDirectGeneration returns allowed for "
    "isShip]",

    "Applies to all of the above: normal authentication and vessel access are separate from these rules and are "
    "always required. These are the conditions the application applies. How strictly the Sail Admin condition binds "
    "depends on the caller's role reaching the server: the 'Generate Now' gate tests the role the client forwards in "
    "the 'x-user-role' header, and when no role header reaches the server it falls back to the built-in identity, "
    "which reports 'Sail Admin'. So the role condition restricts only callers whose client forwards a role. This is "
    "read from the source code; it has not been measured against a running installation. [code: "
    "workOrderGenerationGate.resolveGateRole = forwardedRole || user.role; middleware/auth.ts sets "
    "req.user.role = 'Sail Admin' and stores the forwarded x-user-role separately as forwardedRole]",
]


def set_text(p, text):
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def main():
    d = docx.Document(SRC)
    ps = d.paragraphs
    head = next(i for i, p in enumerate(ps) if p.text.startswith("1.1.14.6 "))
    end = next(i for i, p in enumerate(ps) if p.text.startswith("1.1.14.7 "))
    old = ps[head + 1:end]
    assert all(p.style.name == "List Bullet" for p in old), [p.style.name for p in old]
    set_text(ps[head], HEADING)

    # reuse the first existing bullet for bullet 1, then clone it for the rest so numbering/style is identical
    set_text(old[0], BULLETS[0])
    prev = old[0]
    for text in BULLETS[1:]:
        new = copy.deepcopy(old[0]._element)
        prev._element.addnext(new)
        para = docx.text.paragraph.Paragraph(new, old[0]._parent)
        set_text(para, text)
        prev = para
    for p in old[1:]:                       # drop the remaining original bullets
        p._element.getparent().remove(p._element)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    d.save(OUT)
    for i, p in enumerate(docx.Document(OUT).paragraphs[head - 2:head + 10], head - 2):
        print(f"[{i}|{p.style.name}] {p.text[:110]}")


main()
