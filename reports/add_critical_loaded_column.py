"""Add critical-item SAIL upload counts and rename client status choices."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Font, PatternFill


SOURCE = Path("reports/SAIL_PMS_Progress_Report_Revised.xlsx")
OUTPUT = Path("reports/SAIL_PMS_Progress_Report_Final.xlsx")
STATUS_RENAMES = {
    "Pending Client Input": "Pending Wah Kwong Input",
    "Pending Client Action": "Pending Wah Kwong Action",
}


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    wb = load_workbook(SOURCE)
    track = wb["Update Tracker"]
    dash = wb["Dashboard"]

    # The expected baseline remains hidden. The new loaded count is unknown
    # until confirmed, so it starts blank rather than being assumed zero.
    critical_rows = {}
    for row in range(22, 25):
        critical_rows[row] = {
            "category": track[f"A{row}"].value,
            "received": track[f"B{row}"].value,
            "status": track[f"E{row}"].value,
            "owner": track[f"F{row}"].value,
            "updated": track[f"G{row}"].value,
            "target": track[f"H{row}"].value,
            "remarks": track[f"I{row}"].value,
            "baseline": track[f"N{row}"].value,
            "styles": {
                col: copy(track[f"{col}{row}"]._style)
                for col in "ABCDEFGHIJKL"
            },
        }
    header_styles = {
        col: copy(track[f"{col}21"]._style) for col in "ABCDEFGHIJKL"
    }
    for address in ["I21:L21", "I22:L22", "I23:L23", "I24:L24"]:
        if address in {str(rng) for rng in track.merged_cells.ranges}:
            track.unmerge_cells(address)

    headers = [
        "Critical category",
        "Received from WK",
        "Loaded in SAIL",
        "Pending from WK (est.)",
        "Pending for SAIL upload",
        "Received %",
        "Loaded %",
        "WK verification status",
        "WK owner",
        "Last updated",
        "Target date",
        "Remarks / evidence",
    ]
    style_map = {
        1: "A", 2: "B", 3: "B", 4: "C", 5: "C", 6: "D",
        7: "D", 8: "E", 9: "F", 10: "G", 11: "H", 12: "I",
    }
    for col, header in enumerate(headers, start=1):
        source_col = style_map[col]
        track.cell(21, col)._style = header_styles[source_col]
        track.cell(21, col).value = header
    for row, data in critical_rows.items():
        track[f"N{row}"] = data["baseline"]
        for col, source_col in style_map.items():
            track.cell(row, col)._style = data["styles"][source_col]
        track.cell(row, 1).value = data["category"]
        track.cell(row, 2).value = data["received"]
        track.cell(row, 3).value = None
        track.cell(row, 4).value = (
            f'=IF(AND(ISNUMBER(N{row}),ISNUMBER(B{row}),B{row}<=N{row}),'
            f'N{row}-B{row},"")'
        )
        track.cell(row, 5).value = (
            f'=IF(AND(ISNUMBER(B{row}),ISNUMBER(C{row}),C{row}<=B{row}),'
            f'B{row}-C{row},"")'
        )
        track.cell(row, 6).value = (
            f'=IF(AND(ISNUMBER(B{row}),ISNUMBER(N{row}),N{row}>0,'
            f'B{row}<=N{row}),B{row}/N{row},"")'
        )
        track.cell(row, 7).value = (
            f'=IF(AND(ISNUMBER(C{row}),ISNUMBER(B{row}),ISNUMBER(N{row}),'
            f'N{row}>0,C{row}<=B{row},C{row}<=N{row}),C{row}/N{row},"")'
        )
        track.cell(row, 8).value = data["status"]
        track.cell(row, 9).value = "Wah Kwong / WK Team"
        track.cell(row, 10).value = data["updated"]
        track.cell(row, 11).value = data["target"]
        track.cell(row, 12).value = data["remarks"]
    track["A20"] = (
        "03  /  CLIENT-SIDE — CRITICAL DATA RECEIPT, SAIL LOAD & WK VERIFICATION"
    )
    track["A25"] = (
        "Critical receipt and SAIL loading are tracked separately. Loaded counts "
        "are blank until confirmed; WK verification and outstanding critical spare data remain pending."
    )
    track.row_dimensions[21].height = 48
    for col, width in {
        "A": 26, "B": 16, "C": 16, "D": 18, "E": 19, "F": 13,
        "G": 13, "H": 22, "I": 20, "J": 16, "K": 16, "L": 32,
    }.items():
        track.column_dimensions[col].width = width

    # Preserve the user's counts and use a whole-number dropdown for new loads.
    range_replacements = {
        "B22:B24": "B22:C24",
        "E22:E24": "H22:H24",
        "G22:H24": "J22:K24",
    }
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    if cell.value in STATUS_RENAMES:
                        cell.value = STATUS_RENAMES[cell.value]
                    elif cell.value.startswith("="):
                        for old, new in STATUS_RENAMES.items():
                            cell.value = cell.value.replace(old, new)
        for validation in sheet.data_validations.dataValidation:
            sqref = str(validation.sqref)
            for old, new in range_replacements.items():
                sqref = sqref.replace(old, new)
            validation.sqref = sqref
            if isinstance(validation.formula1, str):
                for old, new in STATUS_RENAMES.items():
                    validation.formula1 = validation.formula1.replace(old, new)

    # Update the dashboard's critical-items chart helper range.
    for helper_row, tracker_row in zip(range(2, 5), range(22, 25)):
        dash[f"Y{helper_row}"] = (
            f'=IF(ISNUMBER(\'Update Tracker\'!B{tracker_row}),'
            f'\'Update Tracker\'!B{tracker_row},"")'
        )
        dash[f"Z{helper_row}"] = (
            f'=IF(ISNUMBER(\'Update Tracker\'!D{tracker_row}),'
            f'\'Update Tracker\'!D{tracker_row},"")'
        )
    dash["J37"] = (
        "Critical chart compares received and pending WK quantities. "
        "SAIL loaded counts are shown in the tracker."
    )

    # Rebuild formatting rules so the new loaded percentage is visible.
    track.conditional_formatting._cf_rules.clear()
    warning_fill = PatternFill(fill_type="solid", fgColor="FADBD8")
    warning_font = Font(color="9C2820")
    for area, formula in [
        ("B7:B11", "AND(ISNUMBER(B7),ISNUMBER(M7),B7>M7)"),
        ("C7:C11", "AND(ISNUMBER(C7),ISNUMBER(B7),C7>B7)"),
        ("B22:B24", "AND(ISNUMBER(B22),ISNUMBER(N22),B22>N22)"),
        ("C22:C24", "AND(ISNUMBER(C22),ISNUMBER(B22),C22>B22)"),
        ("C29:C40", "AND(ISNUMBER(C29),ISNUMBER(B29),C29>B29)"),
        ("C56", "AND(ISNUMBER(C56),ISNUMBER(B56),C56>B56)"),
    ]:
        track.conditional_formatting.add(
            area, FormulaRule(formula=[formula], fill=warning_fill, font=warning_font)
        )
    for area in ["E7:F11", "F22:G24", "E56"]:
        track.conditional_formatting.add(
            area,
            DataBarRule(
                start_type="num", start_value=0, end_type="num", end_value=1,
                color="FF4BA6AA",
            ),
        )

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    track.print_area = "A1:L60"
    dash.print_area = "A1:R40"
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()