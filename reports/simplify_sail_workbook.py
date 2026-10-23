"""Improve readability and repair hidden calculation cells in the SAIL workbook."""

from copy import copy
from pathlib import Path
import re
import warnings
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


SOURCE = Path(
    "attached_assets/SAIL_PMS_Progress_Report_Final_1790769371206.xlsx"
)
OUTPUT = Path("reports/SAIL_PMS_Progress_Report_User_Friendly.xlsx")
MAIN_BASELINE_ROWS = range(7, 12)
CRITICAL_BASELINE_ROWS = range(22, 25)
MAIN_HELPER_HEADER = "Estimated baseline (hidden)"
CRITICAL_HELPER_HEADER = "Estimated critical baseline (hidden)"

MAIN_HEADERS = {
    "A6": "Category",
    "D6": "Pending from WK*",
    "G6": "SAIL status",
    "H6": "Updated",
    "I6": "Target",
    "J6": "Notes",
}
TRACKER_HEADERS = {
    "A15": "Gap",
    "B15": "Components",
    "C15": "Component IDs / names",
    "H15": "Target",
    "I15": "Updated",
    "J15": "Next step / notes",
    "A21": "Critical item",
    "D21": "Pending WK*",
    "E21": "Pending SAIL upload",
    "H21": "WK status",
    "I21": "Updated",
    "J21": "Target",
    "K21": "Notes",
    "A28": "Activity",
    "B28": "Required",
    "C28": "Done",
    "G28": "Target",
    "H28": "Completed",
    "I28": "Updated",
    "J28": "Notes / next action",
    "B45": "WO count",
    "D45": "Updated",
    "E45": "Notes",
    "D55": "RH details pending",
    "G55": "Updated",
    "H55": "Target",
    "I55": "Notes",
}


def _sheet_xml_path(archive: ZipFile, sheet_name: str) -> str:
    ns = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "p": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels.findall("p:Relationship", ns)
    }
    for sheet in workbook.findall("m:sheets/m:sheet", ns):
        if sheet.attrib["name"] == sheet_name:
            target = targets[sheet.attrib[f"{{{ns['r']}}}id"]]
            return str(Path("xl") / target)
    raise KeyError(f"Worksheet not found: {sheet_name}")


def _raw_cell_values(path: Path, sheet_name: str, addresses: list[str]) -> dict:
    """Read cells directly from worksheet XML, including values hidden in merges."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(path) as archive:
        sheet_path = _sheet_xml_path(archive, sheet_name)
        sheet_xml = ET.fromstring(archive.read(sheet_path))
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            strings = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = [
                "".join(text.text or "" for text in item.findall(".//m:t", ns))
                for item in strings.findall("m:si", ns)
            ]
        result = {}
        for address in addresses:
            cell = sheet_xml.find(f'.//m:c[@r="{address}"]', ns)
            if cell is None:
                result[address] = None
                continue
            value = cell.find("m:v", ns)
            if value is None or value.text is None:
                result[address] = None
            elif cell.attrib.get("t") == "s":
                result[address] = shared[int(value.text)]
            elif cell.attrib.get("t") == "b":
                result[address] = value.text == "1"
            else:
                raw = value.text
                try:
                    number = float(raw)
                    result[address] = int(number) if number.is_integer() else number
                except ValueError:
                    result[address] = raw
        return result


def _move_helper_values(source: Path, tracker) -> None:
    addresses = ["M6", "M21"]
    addresses += [f"M{row}" for row in MAIN_BASELINE_ROWS]
    addresses += [f"M{row}" for row in CRITICAL_BASELINE_ROWS]
    values = _raw_cell_values(source, tracker.title, addresses)

    tracker["AA6"] = MAIN_HELPER_HEADER
    tracker["AA6"].comment = Comment(
        "Hidden rough-estimate denominators used by the weighted progress formulas. "
        "Update only when the baseline is confirmed.",
        "Workbook note",
    )
    for row in MAIN_BASELINE_ROWS:
        value = values[f"M{row}"]
        if value is None:
            raise ValueError(f"Missing source estimate at M{row}; refusing to replace it.")
        tracker[f"AA{row}"] = value
    tracker["AA21"] = CRITICAL_HELPER_HEADER
    tracker["AA21"].comment = Comment(
        "Hidden rough-estimate denominators used by the critical-item formulas. "
        "Update only when the baseline is confirmed.",
        "Workbook note",
    )
    for row in CRITICAL_BASELINE_ROWS:
        value = values[f"M{row}"]
        if value is None:
            raise ValueError(f"Missing source estimate at M{row}; refusing to replace it.")
        tracker[f"AA{row}"] = value

    # Remove merges that collided with the original helper cells and reduce
    # full-width section bars to the last visible tracker column.
    for merged in list(tracker.merged_cells.ranges):
        if merged.min_col == 1 and merged.max_col == 12:
            min_row, max_row = merged.min_row, merged.max_row
            tracker.unmerge_cells(str(merged))
            tracker.merge_cells(
                start_row=min_row,
                start_column=1,
                end_row=max_row,
                end_column=11,
            )
        elif merged.min_col >= 12:
            tracker.unmerge_cells(str(merged))

    for row in list(MAIN_BASELINE_ROWS) + list(CRITICAL_BASELINE_ROWS):
        tracker[f"M{row}"] = None
    tracker["M6"] = None
    tracker["M21"] = None

    # Formula references move to a dedicated hidden column beyond the visible
    # report; this avoids future note merges consuming calculation values.
    reference_pattern = re.compile(r"(?<![A-Z0-9_])M(?=\$?\d)")
    for sheet in [tracker, tracker.parent["Dashboard"]]:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = reference_pattern.sub("AA", cell.value)
        for rules in sheet.conditional_formatting._cf_rules.values():
            for rule in rules:
                if rule.formula:
                    rule.formula = [
                        reference_pattern.sub("AA", formula)
                        for formula in rule.formula
                    ]
    for name in tracker.parent.defined_names.values():
        if name.attr_text:
            name.attr_text = reference_pattern.sub("AA", name.attr_text)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    # The source stores denominator values in merged cells. Capture those
    # directly from the XLSX XML before openpyxl normalizes merged cells.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(SOURCE)

    tracker = wb["Update Tracker"]
    dashboard = wb["Dashboard"]
    _move_helper_values(SOURCE, tracker)

    # Keep the report contents and user-entered values; simplify only labels.
    dashboard["A1"] = "SAIL PMS  |  IMPLEMENTATION PROGRESS"
    dashboard["A3"] = (
        "SAIL uploads data received from WK; remaining inputs and verification "
        "are with the Wah Kwong / WK Team."
    )
    dashboard["G8"] = "WK TEAM  /  ACTIONS TO CONFIRM"
    dashboard["J8"] = "WK ACTIVITY STATUS"
    dashboard["B14"] = "Components missing jobs"
    dashboard["B16"] = "Components missing spares"
    dashboard["B18"] = "Critical data received"

    tracker["A1"] = "SAIL PMS  |  UPDATE TRACKER"
    tracker["A3"] = (
        "Enter updates in yellow cells. Blue cells calculate automatically. "
        "Blank means unconfirmed; enter 0 only when confirmed."
    )
    tracker["A4"] = (
        "Quick update: record quantities received from WK and quantities loaded "
        "in SAIL separately; use the dropdowns for statuses."
    )
    tracker["A5"] = "01  /  SAIL TEAM — DATA UPLOAD"
    tracker["A12"] = (
        "Progress uses hidden rough-estimate baselines. Zero-baseline categories "
        "are excluded; Maker list is not weighted."
    )
    tracker["A14"] = "02  /  DATA GAPS — JOBS & SPARES FOR COMPONENTS"
    tracker["A16"] = "Components missing jobs"
    tracker["A17"] = "Components missing spares"
    tracker["A18"] = (
        "List affected components by ID or name. A component with both gaps can "
        "appear in both rows, so do not add the counts."
    )
    tracker["A20"] = "03  /  CRITICAL DATA — RECEIPT, SAIL LOADING & WK CHECK"
    tracker["A25"] = (
        "Track received and loaded separately. Leave unconfirmed quantities blank; "
        "hidden estimates support progress calculations only."
    )
    tracker["A27"] = "04  /  WK TEAM INPUTS, REVIEWS & DECISIONS"
    tracker["A41"] = (
        "Checklist only. Leave unconfirmed quantities, dates, owners and statuses blank."
    )
    tracker["A44"] = "05  /  WORK ORDER SUMMARY — SAIL PMS"
    tracker["A52"] = (
        "Counts may overlap: one job can generate multiple work orders over several cycles."
    )
    tracker["A54"] = "06  /  VESSEL ACTION — RUNNING-HOURS DETAILS"
    tracker["A57"] = (
        "Created = RH machinery in scope. Updated = records with vessel-entered hours. "
        "Pending and percentage calculate from confirmed counts."
    )
    tracker["A40"] = "PMS changes: WK decisions"
    tracker["A59"] = None
    tracker["A60"] = (
        "Set reporting details on Dashboard. Confirm quantities with source owners "
        "before sharing; this workbook is not connected to live SAIL PMS data."
    )

    for address, label in MAIN_HEADERS.items():
        tracker[address] = label
    for address, label in TRACKER_HEADERS.items():
        tracker[address] = label
    tracker["A6"] = "Category"
    tracker["D6"] = "Pending WK*"

    # Improve tracker reading and entry without changing the existing palette.
    if "A4:K4" not in {str(rng) for rng in tracker.merged_cells.ranges}:
        tracker.merge_cells("A4:K4")
    if "A59:K59" in {str(rng) for rng in tracker.merged_cells.ranges}:
        tracker.unmerge_cells("A59:K59")
    tracker["A4"].alignment = Alignment(vertical="center", wrap_text=True)
    tracker.row_dimensions[4].height = 30
    tracker.row_dimensions[59].height = 5
    tracker.row_dimensions[6].height = 42
    tracker.row_dimensions[15].height = 38
    tracker.row_dimensions[21].height = 46
    tracker.row_dimensions[28].height = 40
    tracker.row_dimensions[45].height = 36
    tracker.row_dimensions[55].height = 42
    tracker.freeze_panes = "B7"
    tracker.sheet_view.zoomScale = 85
    tracker.sheet_view.showGridLines = False
    dashboard.sheet_view.showGridLines = False
    dashboard.sheet_view.zoomScale = 90

    for col, width in {
        "A": 34, "B": 16, "C": 17, "D": 18, "E": 16, "F": 14,
        "G": 22, "H": 18, "I": 18, "J": 20, "K": 32,
    }.items():
        tracker.column_dimensions[col].width = width
    for index in range(12, 27):
        tracker.column_dimensions[get_column_letter(index)].hidden = True
    tracker.column_dimensions["AA"].hidden = True
    tracker.column_dimensions["AA"].width = 16
    for index in range(21, 27):
        dashboard.column_dimensions[get_column_letter(index)].hidden = True

    # Wrap section headers and longer user-facing content consistently.
    for row in tracker.iter_rows(min_row=1, max_row=60, min_col=1, max_col=11):
        for cell in row:
            if cell.value is not None:
                alignment = copy(cell.alignment)
                alignment.vertical = "center"
                if alignment.wrap_text is None and cell.row in {
                    3, 4, 6, 12, 15, 18, 21, 25, 28, 41, 45, 52, 55, 57, 60
                }:
                    alignment.wrap_text = True
                cell.alignment = alignment

    # Align printed pages to the visible report width and set the workbook
    # opening view to the stakeholder dashboard.
    tracker.print_area = "A1:K60"
    tracker.page_setup.orientation = "landscape"
    tracker.page_setup.fitToWidth = 1
    tracker.page_setup.fitToHeight = 0
    tracker.sheet_properties.pageSetUpPr.fitToPage = True
    tracker.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.2
    )
    dashboard.print_area = "A1:R40"
    dashboard.page_setup.orientation = "landscape"
    dashboard.page_setup.fitToWidth = 1
    dashboard.page_setup.fitToHeight = 1
    dashboard.sheet_properties.pageSetUpPr.fitToPage = True
    dashboard.page_margins = PageMargins(
        left=0.35, right=0.35, top=0.5, bottom=0.5, header=0.2, footer=0.2
    )
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()